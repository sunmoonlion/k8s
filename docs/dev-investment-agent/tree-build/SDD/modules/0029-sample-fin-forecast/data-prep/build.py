"""把抽出的数写成 Excel（四张工作表 + 出处），并自查。

用法：python -I build.py <extracted.json> <输出 xlsx>
"""

from __future__ import annotations

import json
import re
import sys
from decimal import Decimal

from openpyxl import Workbook

NOTE = re.compile(r"（[^（）]*填列）|\([^()]*填列\)")
KEY_ROWS = [
    "营业收入",
    "归属于上市公司股东的净利润",
    "归属于上市公司股东的扣除非经常性损益的净利润",
    "经营活动产生的现金流量净额",
    "加权平均净资产收益率",
]
KOUFEI = "扣除非经常性损益后的归属于母公司股东的净利润"


def dec(text: str | None) -> Decimal | None:
    if text is None:
        return None
    t = text.replace(",", "")
    neg = t.startswith("(") and t.endswith(")")
    t = t.strip("()")
    if t in ("", "-"):
        return None
    v = Decimal(t)
    return -v if neg else v


def label_of(label: str) -> str:
    return NOTE.sub("", label).strip()


def main() -> None:
    reports = json.load(open(sys.argv[1]))
    periods = [r["period"] for r in reports]
    sheets: dict[str, dict[str, dict[str, Decimal | None]]] = {}
    order: dict[str, list[str]] = {}
    for r in reversed(reports):  # 新的报告定行序，旧报告多出来的行补在后面
        for sheet, rows in r["statements"].items():
            seen: dict[str, int] = {}
            order.setdefault(sheet, [])
            for label, _orig, val in rows:
                name = label_of(label)
                seen[name] = seen.get(name, 0) + 1
                if seen[name] > 1:
                    name = f"{name}#{seen[name]}"
                if name not in order[sheet]:
                    order[sheet].append(name)
                sheets.setdefault(sheet, {}).setdefault(name, {})[r["period"]] = dec(val)
    years = [p for p in periods if p.endswith("12-31")]
    key = {r["period"]: r["key_figures"] for r in reports if r["period"] in years}
    # 扣非：年报「主要会计数据」里的那一项，抄进利润表当一行（报表里没有这一行）
    order["利润表"].append(KOUFEI)
    sheets["利润表"][KOUFEI] = {y: dec(key[y].get(KEY_ROWS[2])) for y in years}

    wb = Workbook()
    wb.remove(wb.active)
    for sheet in ("资产负债表", "利润表", "现金流量表"):
        ws = wb.create_sheet(sheet)
        ws.append(["科目", *periods])
        for name in order[sheet]:
            row = [sheets[sheet][name].get(p) for p in periods]
            if all(v is None for v in row):
                continue  # 每一期都空的行（标题行、本来没有的科目）不要
            ws.append([name, *[float(v) if v is not None else None for v in row]])
    ws = wb.create_sheet("主要会计数据")
    ws.append(["科目", *years])
    for name in KEY_ROWS:
        vals = []
        for y in years:
            v = dec(key[y].get(name))
            if name == "加权平均净资产收益率" and v is not None:
                v = v / 100  # 年报写的是百分数，这里写成小数（带 % 的格子不当数）
            vals.append(float(v) if v is not None else None)
        ws.append([name, *vals])
    ws = wb.create_sheet("出处")
    ws.append(["报告期", "报告", "原件（巨潮资讯网）", "合并资产负债表页", "合并利润表页", "合并现金流量表页", "主要会计数据页"])
    for r in reports:
        p = r["pages"]
        ws.append([r["period"], r["title"], r["url"], p.get("资产负债表"), p.get("利润表"), p.get("现金流量表"),
                   r["key_page"] if r["period"] in years else None])
    ws.append([])
    ws.append(["说明", "合并报表、按报告期累计数（一季报 1–3 月、半年报 1–6 月、三季报 1–9 月、年报全年），单位元；资产负债表为期末时点数。"
               "每期只取报告里「本期」那一栏。扣非净利润与主要会计数据取自各年年报「主要会计数据和财务指标」；加权平均净资产收益率写成小数。"
               "科目名照原文，只去掉「（…填列）」这类填表说明；同一张表里重名的科目按出现次序加 #2。"])
    wb.save(sys.argv[2])

    # 自查：四条勾稽、报表与主要会计数据
    bs, inc, cf = sheets["资产负债表"], sheets["利润表"], sheets["现金流量表"]
    def g(t, name, p):
        return t.get(name, {}).get(p)
    problems = []
    for p in periods:
        a, l, e = g(bs, "资产总计", p), g(bs, "负债合计", p), g(bs, "所有者权益（或股东权益）合计", p)
        if None in (a, l, e) or abs(a - (l + e)) > 1:
            problems.append(("bs", p, a, l, e))
        pa, m, n = g(inc, "归属于母公司股东的净利润", p), g(inc, "少数股东损益", p), g(inc, "净利润", p)
        if None in (pa, m, n) or abs(pa + m - n) > 1:
            problems.append(("np", p, pa, m, n))
        e2, b, d = g(cf, "期末现金及现金等价物余额", p), g(cf, "期初现金及现金等价物余额", p), g(cf, "现金及现金等价物净增加额", p)
        if None in (e2, b, d) or abs(e2 - (b + d)) > 1:
            problems.append(("cash", p, e2, b, d))
    for prev, cur in zip(years, years[1:]):
        o, c = g(cf, "期初现金及现金等价物余额", cur), g(cf, "期末现金及现金等价物余额", prev)
        if None in (o, c) or abs(o - c) > 1:
            problems.append(("opening", cur, o, c))
    for y in years:
        for ours, kname in (("营业收入", KEY_ROWS[0]), ("归属于母公司股东的净利润", KEY_ROWS[1])):
            if g(inc, ours, y) != dec(key[y].get(kname)):
                problems.append(("key", y, ours, g(inc, ours, y), key[y].get(kname)))
        if g(cf, "经营活动产生的现金流量净额", y) != dec(key[y].get(KEY_ROWS[3])):
            problems.append(("key", y, "经营现金流", g(cf, "经营活动产生的现金流量净额", y), key[y].get(KEY_ROWS[3])))
    print("rows:", {s: len(order[s]) for s in order}, "periods:", len(periods))
    print("problems:", len(problems))
    for x in problems:
        print("  ", x)


if __name__ == "__main__":
    main()
